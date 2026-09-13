from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUT = 'TrackShift_Model_Postulates_and_Graph_Insights_Report.docx'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9D9D9', size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    trPr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color='000000', size=9.0, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def remove_paragraph_borders(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is not None:
        pPr.remove(pBdr)


def add_table(doc, headers, rows, widths, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        hdr.cells[i].width = Inches(widths[i])
        set_cell_shading(hdr.cells[i], '17365D')
        set_cell_text(hdr.cells[i], h, bold=True, color='FFFFFF', size=font_size)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_shading(cells[i], 'F2F6FA' if ridx % 2 else 'FFFFFF')
            set_cell_text(cells[i], value, size=font_size)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_para(doc, text='', style=None, bold_lead=None, space_after=4, size=9.8):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.06
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        r1.bold = True
        r1.font.name = 'Arial'
        r1._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
        r1._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
        r1.font.size = Pt(size)
        r2 = p.add_run(text[len(bold_lead):])
        r2.font.name = 'Arial'
        r2._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
        r2._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
        r2.font.size = Pt(size)
    else:
        r = p.add_run(text)
        r.font.name = 'Arial'
        r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
        r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
        r.font.size = Pt(size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.bold = True
    r.font.size = Pt(13 if level == 1 else 10.5)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run('Page ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
normal.font.size = Pt(9.8)
normal.font.color.rgb = RGBColor(0, 0, 0)
for name, size in [('Title', 20), ('Heading 1', 13), ('Heading 2', 10.5)]:
    st = styles[name]
    st.font.name = 'Arial'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(size)
    st.font.bold = True
    pPr = st._element.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is not None:
        pPr.remove(pBdr)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_before = Pt(0)
fr = fp.add_run('TrackShift F1 HEV Prototype | Interpretation Report  |  ')
fr.font.name = 'Arial'
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(100, 100, 100)
add_page_number(fp)

title = doc.add_paragraph(style='Title')
remove_paragraph_borders(title)
title.paragraph_format.space_after = Pt(2)
title.paragraph_format.keep_with_next = True
tr = title.add_run('TrackShift F1 HEV Model Postulates and Graph Insights')
tr.font.name = 'Arial'
tr._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
tr._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
tr.font.size = Pt(20)
tr.font.color.rgb = RGBColor(0, 0, 0)
tr.bold = True

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(8)
sr = sub.add_run('Technical interpretation of the graphical Simulink model and detailed MATLAB analysis')
sr.font.name = 'Arial'
sr.font.size = Pt(10.5)
sr.font.italic = True
sr.font.color.rgb = RGBColor(80, 80, 80)

add_para(doc, 'Purpose and conclusion. This report translates the prototype outputs into engineering statements that can be tested, explained to judges and used to guide the next development stage. The model supports a clear conclusion: selective MGU-K deployment can be evaluated against a no-overtake counterfactual, but the current result is a system-level demonstrator rather than a calibrated F1 power-unit prediction. Its strongest contribution is the traceable link from race context to AI recommendation, constrained electrical power, battery state and vehicle response.', space_after=5)

add_heading(doc, '1 Model operation in one page', 1)
add_para(doc, 'The model runs two parallel branches from identical speed, gap, throttle, brake, track-position and gradient signals. The upper branch has overtake permission enabled; the lower branch has deployment disabled but still permits regenerative braking. The AI estimates pass probability, rival-cover probability and expected gain. A rule gate checks eligibility, after which the power-command governor applies battery, torque, speed, physical and power limits. The resulting MGU-K command passes through the inverter, MGU-K, gearbox and vehicle dynamics, while the battery state is updated through a discrete energy memory.', space_after=4)
add_para(doc, 'The electromechanical interpretation is: fuel produces ICE power; the battery can provide additional MGU-K mechanical power through the inverter; braking can reverse the MGU-K power flow and recharge the battery. Positive MGU-K power is deployment, negative power is harvesting and zero power is hold. The MATLAB analysis adds detailed logs, scenario sweeps, validation tests and an efficiency study.', space_after=4)

add_heading(doc, '2 Verified prototype behaviour', 1)
add_table(doc,
    ['Observed behaviour', 'What it demonstrates', 'Engineering interpretation'],
    [
        ['AI branch reaches +275 kW MGU-K deployment', 'The overtake-enabled branch can spend electrical energy when its conditions are met.', 'The AI is connected to an actionable power request, not only a display.'],
        ['No-overtake branch has no positive MGU-K deployment', 'The counterfactual branch blocks overtake assistance.', 'The comparison isolates the value of the AI-enabled policy.'],
        ['Both branches can show negative MGU-K power', 'Regeneration remains available during braking.', 'No-overtake means no deployment, not no energy recovery.'],
        ['Battery energy changes and SOC remains bounded', 'The energy store is a state variable with limits.', 'The governor and BMS prevent unrestricted battery use.'],
        ['Seven MATLAB validation tests pass', 'Hold, harvest, limit, torque, swing, recharge and eligibility checks run automatically.', 'The prototype has internal consistency checks before physical calibration.'],
    ], [2.0, 2.35, 2.85], font_size=8.4)

doc.add_page_break()
add_heading(doc, '3 Postulates derived from the model', 1)
add_para(doc, 'A postulate here is a model-supported engineering hypothesis, not a proven real-world fact. Each postulate is useful because it suggests a measurable test or design action for the next version.', space_after=5)
add_table(doc,
    ['Postulate', 'Evidence in the model', 'Implication for improvement'],
    [
        ['Selective deployment should deliver better energy value than blanket deployment.', 'Deployment is permitted only when expected gain exceeds the threshold; weak opportunities are held.', 'Compare time gained per MJ spent across many overtaking zones, not only total lap time.'],
        ['The AI recommendation is not sufficient by itself; the governor determines the physically usable action.', 'The requested power is passed through regulatory, speed, physical and battery limits before becoming final power.', 'Log requested, final, clipped and limiting-reason signals to identify the active bottleneck.'],
        ['Regeneration reduces the net energy cost of deployment but does not make an overtake energy-free.', 'The battery charges under negative MGU-K power and loses energy under positive power; conversion losses are included.', 'Report gross deployment energy, recovered energy and net battery energy separately.'],
        ['Race context is as important as power capability.', 'Gap and track position affect the eligibility gate; the same powertrain behaves differently in different zones.', 'Use real circuit zones and measured relative-speed data instead of a generic track-position window.'],
        ['A no-overtake branch is a counterfactual baseline, not an alternative AI opinion.', 'Both branches calculate similar predictions from the same inputs, but only the upper branch can act on the recommendation.', 'Explain policy permission and AI decision as separate signals in demonstrations and analysis.'],
        ['Battery operating margin is a strategic resource.', 'The BMS and governor restrict the battery to an operating window and energy swing.', 'Add multi-lap energy planning so the controller protects energy for later, higher-value zones.'],
        ['Power-electronics efficiency materially affects strategy quality.', 'The efficiency sweep changes energy use and harvesting while the decision logic remains the same.', 'Prioritise inverter, MGU-K and thermal-loss calibration when ranking hardware improvements.'],
        ['The current model is strongest as a traceable control demonstrator, not as a final vehicle predictor.', 'Synthetic telemetry and simplified ICE, battery, turbo and vehicle equations are used.', 'Replace replay inputs with measured data and close the speed/driver feedback loop before making performance claims.'],
    ], [2.12, 2.52, 2.56], font_size=8.0)

doc.add_page_break()
add_heading(doc, '4 Graph by graph understanding and insight', 1)
add_para(doc, 'The dashboard should be read as a causal chain: eligibility creates an AI decision, the governor converts that decision into permitted power, the power changes the energy store, and the vehicle response shows the downstream effect.', space_after=4)
add_table(doc,
    ['Graph or display', 'Understanding from the current prototype', 'Insight and next action'],
    [
        ['Speed response', 'Compares baseline, TrackShift and source speed. Similar curves show that the replay is closely followed.', 'Useful for checking signal consistency, but not enough to prove lap-time improvement. Add closed-loop vehicle speed and a driver model.'],
        ['Governed power', 'Shows final MGU-K command against regulatory and absolute limits. Spikes are deployment; negative regions are harvest.', 'The gap between requested and final power identifies clipping. Add an explicit active-limit label to make the cause visible.'],
        ['Energy Store state', 'Downward movement means battery discharge; upward movement means recovered energy. The initial-energy line gives the reference.', 'Measure energy spent per deployment event and energy recovered after it. A final SOC alone can hide the cost of earlier deployment.'],
        ['Eligibility context', 'Shows gap against the detection threshold and when overtake activation is present.', 'Tests whether activation occurs in a plausible race context. Replace synthetic gap profiles with measured rival trajectories.'],
        ['Decision engine', 'Shows AI score, pass probability and the deployment threshold. Crossing the threshold supports a recommendation, not a guaranteed pass.', 'Evaluate calibration using real labels, precision/recall and Brier score; do not judge the AI only by visual threshold crossings.'],
        ['Controller action', 'Shows deploy, hold and harvest states alongside governor clipping.', 'Correlate every decision transition with MGU-K power and brake input. Frequent clipping indicates the AI request or threshold is not aligned with hardware limits.'],
        ['MGU-K power Scope', 'The AI branch has a positive deployment pulse; the no-overtake branch remains at zero for positive deployment. Negative sections are regeneration.', 'This is the clearest proof of mode separation. Report pulse duration, peak power and integrated energy.'],
        ['Battery energy Scope', 'Compares how the two modes spend and recover energy over time.', 'A useful strategy should trade a controlled energy decrease for measurable performance benefit, then preserve enough margin for later zones.'],
    ], [1.55, 2.72, 2.93], font_size=7.9)

doc.add_page_break()
add_heading(doc, '5 What the graphs say about solution quality', 1)
add_para(doc, 'The graphs collectively indicate that the prototype has a coherent control chain. A positive MGU-K pulse appears only when the overtake-enabled policy permits it, while the no-overtake branch remains a valid hold/regen comparison. Battery energy responds in the physically expected direction, and the governor keeps the command within configured limits. These are strong demonstrations of connectivity and control logic.', space_after=4)
add_para(doc, 'The graphs do not yet prove that the AI is faster or more accurate than a real race engineer. The current replay is synthetic, the powertrain maps are simplified, and the displayed final values are instantaneous. The strongest evidence is therefore consistency of cause and effect: a change in race context should change eligibility; eligibility should change the AI action; the action should change MGU-K power; and the power should change battery energy and vehicle response.', space_after=4)

add_heading(doc, '6 Recommended improvements', 1)
add_table(doc,
    ['Priority', 'Improvement', 'Why it matters'],
    [
        ['1', 'Replace synthetic telemetry with measured or validated open telemetry.', 'Makes gap, speed, braking and track-zone conclusions representative of real use.'],
        ['2', 'Close the vehicle-speed and driver feedback loop.', 'Allows the model to predict the effect of deployment instead of replaying a prescribed speed trace.'],
        ['3', 'Add full time-history logging for all branch outputs.', 'Makes engine, battery, BMS, torque, forces and clipping directly exportable from Simulink.'],
        ['4', 'Calibrate ICE, turbo, MGU-K, inverter, battery voltage-current and thermal maps.', 'Converts the architecture demonstrator into a defensible engineering performance model.'],
        ['5', 'Add multi-lap energy planning and zone-specific deployment rules.', 'Prevents spending energy in a low-value zone when a later overtake has higher expected value.'],
        ['6', 'Validate AI predictions using labelled overtaking outcomes.', 'Separates a good-looking score trace from a statistically reliable decision model.'],
    ], [0.55, 3.55, 3.10], font_size=8.2)

add_heading(doc, '7 Scope and limitations', 1)
add_para(doc, 'The model represents the major system boundaries: ICE, turbocharger, fuel delivery, MGU-K, inverter, battery, BMS, braking, gearbox, aerodynamics, tyres and longitudinal dynamics. It is a native Simulink signal-level model and a pure-MATLAB reference implementation. It is not a confidential F1 team calibration, a detailed electrochemical battery model, a combustion CFD model, a hardware-in-the-loop model or a regulatory certification model.', space_after=4)
add_para(doc, 'The correct presentation claim is: “This prototype demonstrates how AI overtake intelligence can be connected to a constrained hybrid powertrain and evaluated through battery, MGU-K, engine and vehicle metrics.” The next credible step is calibration and closed-loop validation, not simply increasing the number of blocks in the diagram.', space_after=4)

add_heading(doc, '8 Final conclusion', 1)
add_para(doc, 'The model supports the postulate that energy should be deployed selectively, only when the predicted strategic value is high enough and the battery, machine and regulatory constraints allow it. The most informative evidence is the combined reading of governed MGU-K power, battery-energy trajectory, eligibility context and controller action. Together these graphs reveal not only what the controller did, but why it did it and whether the powertrain could safely execute the decision.', space_after=2)

doc.core_properties.title = 'TrackShift F1 HEV Model Postulates and Graph Insights'
doc.core_properties.subject = 'Engineering interpretation of the TrackShift HEV prototype'
doc.core_properties.author = 'TrackShift Project'
doc.core_properties.comments = 'Generated technical report'
doc.save(OUT)
print(OUT)
